"""
Build the user guides (Markdown) into Word (.docx) and PDF files.

    python docs/build_docs.py

Reads docs/HUONG_DAN_*.md and writes .docx and .pdf next to them. Needs:
python-docx, playwright (with Chromium installed: `playwright install chromium`).
Only the small Markdown subset used by the guides is supported: #/##/### headings,
paragraphs, bullet/numbered lists (nested by indentation), tables, > quotes, fenced
code, --- rules and **bold** / *italic* / `code` / [links](#anchor or url).
"""
import html
import re
import sys
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).resolve().parent
GUIDES = [
    ("HUONG_DAN_ADMIN", "Hướng dẫn sử dụng – Quản trị (Admin)"),
    ("HUONG_DAN_KHACH_HANG", "Hướng dẫn sử dụng – Đối tác / Nhà phân phối"),
]
BRAND = "1E9A92"        # headings, links accents
BRAND_TINT = "E6F6F5"   # table header / quote background
GRAY_BORDER = "BFCBD5"
CODE_BG = "F1F5F9"
BODY_FONT = "Calibri"
CODE_FONT = "Consolas"


# ───────────────────────── Markdown → AST ─────────────────────────
def slug(text: str) -> str:
    t = unicodedata.normalize("NFC", text).lower()
    return re.sub(r"[^\w\- ]", "", t).replace(" ", "-")


INLINE_RE = re.compile(
    r"(`[^`]+`)|(\[([^\]]+)\]\(([^)]+)\))|(\*\*(.+?)\*\*)|(\*(?!\s)(.+?)(?<!\s)\*)"
)


def parse_inline(s: str) -> list:
    out, pos = [], 0
    for m in INLINE_RE.finditer(s):
        if m.start() > pos:
            out.append(("text", s[pos:m.start()]))
        if m.group(1):
            out.append(("code", m.group(1)[1:-1]))
        elif m.group(2):
            out.append(("link", parse_inline(m.group(3)), m.group(4)))
        elif m.group(5):
            out.append(("bold", parse_inline(m.group(6))))
        else:
            out.append(("italic", parse_inline(m.group(8))))
        pos = m.end()
    if pos < len(s):
        out.append(("text", s[pos:]))
    return out


def plain(nodes) -> str:
    return "".join(
        n[1] if n[0] in ("text", "code") else plain(n[1]) for n in nodes
    )


LIST_RE = re.compile(r"^(\s*)([-*]|\d+\.)\s+(.*)$")


def indent_of(line: str) -> int:
    return len(line) - len(line.lstrip(" "))


def parse_list(lines, i, base):
    """Parse a list whose items sit at indent `base`. Returns (block, next_index)."""
    m = LIST_RE.match(lines[i])
    ordered = m.group(2)[0].isdigit()
    items = []
    n = len(lines)
    while i < n:
        line = lines[i]
        if not line.strip():
            j = i
            while j < n and not lines[j].strip():
                j += 1
            if j < n and LIST_RE.match(lines[j]) and indent_of(lines[j]) >= base:
                i = j
                continue
            break
        m = LIST_RE.match(line)
        ind = indent_of(line)
        if m and ind == base:
            if m.group(2)[0].isdigit() != ordered:
                break
            items.append({"text": m.group(3).strip(), "children": []})
            i += 1
        elif m and ind > base and items:
            child, i = parse_list(lines, i, ind)
            items[-1]["children"].append(child)
        elif ind > base and items and line.lstrip().startswith(">"):
            quote = []
            while i < n and lines[i].strip().startswith(">") and indent_of(lines[i]) > base:
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            items[-1]["children"].append({"type": "quote", "text": " ".join(quote)})
        elif ind > base and items:
            items[-1]["text"] += " " + line.strip()
            i += 1
        else:
            break
    return {"type": "ol" if ordered else "ul", "items": items}, i


def split_row(line: str) -> list:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def parse_markdown(text: str) -> list:
    lines = text.splitlines()
    blocks, i, n = [], 0, len(lines)
    while i < n:
        line = lines[i]
        if not line.strip():
            i += 1
        elif line.startswith("```"):
            i += 1
            code = []
            while i < n and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append({"type": "code", "text": "\n".join(code)})
        elif re.match(r"^#{1,3} ", line):
            level = len(line) - len(line.lstrip("#"))
            blocks.append({"type": "h", "level": level, "text": line[level:].strip()})
            i += 1
        elif re.match(r"^-{3,}\s*$", line):
            blocks.append({"type": "hr"})
            i += 1
        elif line.lstrip().startswith("|") and i + 1 < n and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            header = split_row(line)
            i += 2
            rows = []
            while i < n and lines[i].lstrip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            blocks.append({"type": "table", "header": header, "rows": rows})
        elif line.startswith(">"):
            quote = []
            while i < n and lines[i].startswith(">"):
                quote.append(lines[i].lstrip(">").strip())
                i += 1
            blocks.append({"type": "quote", "text": " ".join(quote)})
        elif LIST_RE.match(line):
            block, i = parse_list(lines, i, indent_of(line))
            blocks.append(block)
        else:
            para = [line.strip()]
            i += 1
            while (
                i < n and lines[i].strip() and not lines[i].startswith(("#", ">", "```"))
                and not LIST_RE.match(lines[i]) and not lines[i].lstrip().startswith("|")
                and not re.match(r"^-{3,}\s*$", lines[i])
            ):
                para.append(lines[i].strip())
                i += 1
            blocks.append({"type": "p", "text": " ".join(para)})
    return blocks


# ───────────────────────── HTML / PDF ─────────────────────────
CSS = f"""
@page {{ size: A4; }}
* {{ box-sizing: border-box; }}
body {{ font-family: 'Segoe UI', Calibri, Arial, sans-serif; font-size: 10.5pt; line-height: 1.5; color: #1e293b; }}
h1 {{ font-size: 24pt; color: #{BRAND}; margin: 0 0 6pt; line-height: 1.2; }}
h2 {{ font-size: 16pt; color: #{BRAND}; border-bottom: 2px solid #{BRAND}; padding-bottom: 3pt; margin: 22pt 0 8pt; break-after: avoid; }}
h2.pb {{ break-before: page; margin-top: 0; }}
h3 {{ font-size: 12.5pt; color: #0f172a; margin: 14pt 0 5pt; break-after: avoid; }}
p {{ margin: 0 0 7pt; }}
ul, ol {{ margin: 0 0 7pt; padding-left: 20pt; }}
li {{ margin-bottom: 3pt; }}
a {{ color: #0563c1; text-decoration: none; }}
code {{ font-family: Consolas, 'Courier New', monospace; font-size: 9.5pt; background: #{CODE_BG}; padding: 1px 4px; border-radius: 3px; }}
pre {{ font-family: Consolas, 'Courier New', monospace; font-size: 8.8pt; line-height: 1.35; background: #{CODE_BG}; border: 1px solid #{GRAY_BORDER}; border-radius: 4px; padding: 8pt 10pt; margin: 0 0 9pt; white-space: pre-wrap; break-inside: avoid; }}
pre code {{ background: none; padding: 0; font-size: inherit; }}
blockquote {{ margin: 0 0 8pt; padding: 6pt 10pt; background: #{BRAND_TINT}; border-left: 4px solid #{BRAND}; border-radius: 0 4px 4px 0; break-inside: avoid; }}
li blockquote {{ margin-top: 4pt; }}
hr {{ border: 0; border-top: 1px solid #{GRAY_BORDER}; margin: 12pt 0; }}
table {{ border-collapse: collapse; width: 100%; margin: 0 0 10pt; font-size: 9.6pt; }}
th, td {{ border: 1px solid #{GRAY_BORDER}; padding: 4pt 6pt; vertical-align: top; text-align: left; }}
th {{ background: #{BRAND_TINT}; font-weight: 700; }}
tr {{ break-inside: avoid; }}
thead {{ display: table-header-group; }}
"""


def inline_html(nodes, anchors) -> str:
    out = []
    for n in nodes:
        if n[0] == "text":
            out.append(html.escape(n[1]))
        elif n[0] == "code":
            out.append(f"<code>{html.escape(n[1])}</code>")
        elif n[0] == "bold":
            out.append(f"<strong>{inline_html(n[1], anchors)}</strong>")
        elif n[0] == "italic":
            out.append(f"<em>{inline_html(n[1], anchors)}</em>")
        else:
            url = n[2]
            if url.startswith("#"):
                if url[1:] not in anchors:
                    out.append(inline_html(n[1], anchors))
                    continue
            out.append(f'<a href="{html.escape(url)}">{inline_html(n[1], anchors)}</a>')
    return "".join(out)


def block_html(b, anchors) -> str:
    t = b["type"]
    if t == "h":
        cls = ' class="pb"' if b["level"] == 2 and re.match(r"^1\.\s", b["text"]) else ""
        nodes = parse_inline(b["text"])
        return f'<h{b["level"]}{cls} id="{slug(plain(nodes))}">{inline_html(nodes, anchors)}</h{b["level"]}>'
    if t == "p":
        return f"<p>{inline_html(parse_inline(b['text']), anchors)}</p>"
    if t == "hr":
        return "<hr>"
    if t == "quote":
        return f"<blockquote>{inline_html(parse_inline(b['text']), anchors)}</blockquote>"
    if t == "code":
        return f"<pre><code>{html.escape(b['text'])}</code></pre>"
    if t in ("ul", "ol"):
        lis = []
        for it in b["items"]:
            inner = inline_html(parse_inline(it["text"]), anchors) + "".join(
                block_html(c, anchors) for c in it["children"]
            )
            lis.append(f"<li>{inner}</li>")
        return f"<{t}>{''.join(lis)}</{t}>"
    if t == "table":
        head = "".join(f"<th>{inline_html(parse_inline(c), anchors)}</th>" for c in b["header"])
        rows = "".join(
            "<tr>" + "".join(f"<td>{inline_html(parse_inline(c), anchors)}</td>" for c in r) + "</tr>"
            for r in b["rows"]
        )
        return f"<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>"
    raise ValueError(t)


def build_html(blocks, title, anchors) -> str:
    body = "\n".join(block_html(b, anchors) for b in blocks)
    return f'<!DOCTYPE html><html lang="vi"><head><meta charset="utf-8"><title>{html.escape(title)}</title><style>{CSS}</style></head><body>{body}</body></html>'


def build_pdf(html_text: str, out: Path, header: str):
    from playwright.sync_api import sync_playwright

    tmp = out.with_suffix(".tmp.html")
    tmp.write_text(html_text, encoding="utf-8")
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page()
            page.goto(tmp.as_uri())
            page.pdf(
                path=str(out), format="A4", print_background=True,
                margin={"top": "2.2cm", "bottom": "2cm", "left": "2cm", "right": "2cm"},
                display_header_footer=True,
                header_template=(
                    '<div style="font-family:Segoe UI,Arial,sans-serif;font-size:8px;color:#64748b;'
                    f'width:100%;padding:0 2cm;text-align:right;">{html.escape(header)}</div>'
                ),
                footer_template=(
                    '<div style="font-family:Segoe UI,Arial,sans-serif;font-size:8px;color:#64748b;'
                    'width:100%;text-align:center;">Trang <span class="pageNumber"></span> / '
                    '<span class="totalPages"></span></div>'
                ),
            )
            browser.close()
    finally:
        tmp.unlink(missing_ok=True)


# ───────────────────────── DOCX ─────────────────────────
TEXT_WIDTH_DXA = 9638  # A4 (21cm) minus 2 x 2cm margins


def set_font(style_or_run_font_owner, name: str, rpr):
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in list(rfonts.attrib):
        del rfonts.attrib[attr]
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def style_font(style, name, size=None, bold=None, color=None):
    style.font.name = name
    set_font(style, name, style.element.get_or_add_rPr())
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)


PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
    "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
    "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
    "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
    "sectPr", "pPrChange",
]
TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
    "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar",
    "tblLook", "tblCaption", "tblDescription",
]


def insert_ordered(parent, el, order):
    """Insert `el` into `parent` at its schema position (Word is strict about child order)."""
    name = el.tag.split("}")[1]
    idx = order.index(name)
    for child in parent:
        cname = child.tag.split("}")[1]
        if cname in order and order.index(cname) > idx:
            child.addprevious(el)
            return
    parent.append(el)


def shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    if el_pr.tag == qn("w:pPr"):
        insert_ordered(el_pr, shd, PPR_ORDER)
    else:
        el_pr.append(shd)


def para_border(par, side, color, size=8, space=4):
    ppr = par._p.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        insert_ordered(ppr, bdr, PPR_ORDER)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    bdr.append(el)


def add_field(par, instr, size=None):
    for kind in ("begin", None, "separate", "text", "end"):
        run = par.add_run()
        if size:
            run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        if kind == "text":
            run.text = "1"
        elif kind is None:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instr} "
            run._r.append(it)
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)


class DocxBuilder:
    def __init__(self, title, header_text):
        self.doc = Document()
        self.anchors = {}   # slug -> bookmark name
        self.bm_id = 100
        self._setup(title, header_text)
        self._setup_numbering()

    # -- setup --
    def _setup(self, title, header_text):
        d = self.doc
        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2)
        d.core_properties.title = title
        d.core_properties.author = "IMV"
        d.core_properties.language = "vi-VN"
        zoom = d.settings.element.find(qn("w:zoom"))
        if zoom is not None:
            zoom.set(qn("w:percent"), "100")

        style_font(d.styles["Normal"], BODY_FONT, 11)
        pf = d.styles["Normal"].paragraph_format
        pf.space_after, pf.line_spacing = Pt(6), 1.15
        style_font(d.styles["Heading 1"], BODY_FONT, 26, True, BRAND)
        style_font(d.styles["Heading 2"], BODY_FONT, 16, True, BRAND)
        style_font(d.styles["Heading 3"], BODY_FONT, 13, True, "0F172A")
        for name, before, after in (("Heading 1", 0, 6), ("Heading 2", 20, 8), ("Heading 3", 12, 4)):
            f = d.styles[name].paragraph_format
            f.space_before, f.space_after, f.keep_with_next = Pt(before), Pt(after), True

        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run(header_text)
        r.font.size, r.font.color.rgb = Pt(8), RGBColor(0x64, 0x74, 0x8B)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("Trang ")
        r.font.size, r.font.color.rgb = Pt(8), RGBColor(0x64, 0x74, 0x8B)
        add_field(fp, "PAGE", 8)
        r = fp.add_run(" / ")
        r.font.size, r.font.color.rgb = Pt(8), RGBColor(0x64, 0x74, 0x8B)
        add_field(fp, "NUMPAGES", 8)

    def _setup_numbering(self):
        numbering = self.doc.part.numbering_part.element

        def lvl(i, fmt, text, left, hanging=360):
            return (
                f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
                f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
                f'<w:pPr><w:ind w:left="{left}" w:hanging="{hanging}"/></w:pPr></w:lvl>'
            )

        bullet = (
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="100"><w:multiLevelType w:val="hybridMultilevel"/>'
            + lvl(0, "bullet", "•", 567) + lvl(1, "bullet", "–", 1021) + lvl(2, "bullet", "•", 1474)
            + "</w:abstractNum>"
        )
        decimal = (
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="101"><w:multiLevelType w:val="hybridMultilevel"/>'
            + lvl(0, "decimal", "%1.", 567) + lvl(1, "lowerLetter", "%2.", 1021) + lvl(2, "lowerRoman", "%3.", 1474)
            + "</w:abstractNum>"
        )
        first_num = numbering.find(qn("w:num"))
        for x in (bullet, decimal):
            el = parse_xml(x)
            if first_num is not None:
                first_num.addprevious(el)
            else:
                numbering.append(el)
        self.next_num = 200
        self.bullet_num = self._new_num(100)

    def _new_num(self, abstract_id, restart=False):
        numbering = self.doc.part.numbering_part.element
        nid = self.next_num
        self.next_num += 1
        override = (
            '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride>' if restart else ""
        )
        numbering.append(parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{nid}"><w:abstractNumId w:val="{abstract_id}"/>{override}</w:num>'
        ))
        return nid

    # -- inline --
    def add_inline(self, par, nodes, size=None, bold=False, italic=False, link_wrap=None):
        for n in nodes:
            kind = n[0]
            if kind in ("text", "code"):
                run = par.add_run(n[1])
                run.bold, run.italic = bold or None, italic or None
                if size:
                    run.font.size = Pt(size)
                if kind == "code":
                    rpr = run._r.get_or_add_rPr()
                    set_font(run, CODE_FONT, rpr)
                    run.font.size = Pt((size or 11) - 1)
                    shade(rpr, CODE_BG)
                if link_wrap is not None:
                    run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
                    run.font.underline = True
                    link_wrap.append(run._r)
            elif kind == "bold":
                self.add_inline(par, n[1], size, True, italic, link_wrap)
            elif kind == "italic":
                self.add_inline(par, n[1], size, bold, True, link_wrap)
            else:  # link
                url = n[2]
                anchor = url[1:] if url.startswith("#") else None
                if anchor is not None and anchor not in self.anchors:
                    self.add_inline(par, n[1], size, bold, italic, link_wrap)
                    continue
                h = OxmlElement("w:hyperlink")
                if anchor is not None:
                    h.set(qn("w:anchor"), self.anchors[anchor])
                else:
                    h.set(qn("r:id"), par.part.relate_to(url, RT.HYPERLINK, is_external=True))
                par._p.append(h)
                self.add_inline(par, n[1], size, bold, italic, h)

    # -- blocks --
    def add_block(self, b, list_level=0):
        t = b["type"]
        d = self.doc
        if t == "h":
            nodes = parse_inline(b["text"])
            par = d.add_paragraph(style=f"Heading {b['level']}")
            self.add_inline(par, nodes)
            if b["level"] == 2:
                para_border(par, "bottom", BRAND, 8, 2)
                if re.match(r"^1\.\s", b["text"]):
                    par.paragraph_format.page_break_before = True
                    par.paragraph_format.space_before = Pt(0)
            name = self.anchors.get(slug(plain(nodes)))
            if name:
                self.bm_id += 1
                start = OxmlElement("w:bookmarkStart")
                start.set(qn("w:id"), str(self.bm_id))
                start.set(qn("w:name"), name)
                end = OxmlElement("w:bookmarkEnd")
                end.set(qn("w:id"), str(self.bm_id))
                ppr = par._p.find(qn("w:pPr"))
                (ppr.addnext(start) if ppr is not None else par._p.insert(0, start))
                par._p.append(end)
        elif t == "p":
            par = d.add_paragraph()
            self.add_inline(par, parse_inline(b["text"]))
        elif t == "hr":
            par = d.add_paragraph()
            par.paragraph_format.space_before = par.paragraph_format.space_after = Pt(6)
            para_border(par, "bottom", GRAY_BORDER, 6, 1)
        elif t == "quote":
            self.add_quote(b["text"], 0)
        elif t == "code":
            par = d.add_paragraph()
            ppr = par._p.get_or_add_pPr()
            shade(ppr, CODE_BG)
            for side in ("top", "left", "bottom", "right"):
                para_border(par, side, GRAY_BORDER, 4, 4)
            par.paragraph_format.left_indent = Cm(0.2)
            par.paragraph_format.right_indent = Cm(0.2)
            par.paragraph_format.line_spacing = 1.0
            for i, line in enumerate(b["text"].split("\n")):
                run = par.add_run(line)
                set_font(run, CODE_FONT, run._r.get_or_add_rPr())
                run.font.size = Pt(8.5)
                if i < len(b["text"].split("\n")) - 1:
                    run.add_break()
        elif t in ("ul", "ol"):
            self.add_list(b, 0)
        elif t == "table":
            self.add_table(b)

    def add_quote(self, text, indent_cm):
        par = self.doc.add_paragraph()
        ppr = par._p.get_or_add_pPr()
        shade(ppr, BRAND_TINT)
        para_border(par, "left", BRAND, 24, 6)
        par.paragraph_format.left_indent = Cm(0.4 + indent_cm)
        par.paragraph_format.space_after = Pt(6)
        self.add_inline(par, parse_inline(text))

    def add_list(self, b, level):
        d = self.doc
        num = self.bullet_num if b["type"] == "ul" else self._new_num(101, restart=True)
        for it in b["items"]:
            par = d.add_paragraph()
            par.paragraph_format.space_after = Pt(3)
            ppr = par._p.get_or_add_pPr()
            numpr = OxmlElement("w:numPr")
            il = OxmlElement("w:ilvl")
            il.set(qn("w:val"), str(min(level, 2)))
            ni = OxmlElement("w:numId")
            ni.set(qn("w:val"), str(num))
            numpr.append(il)
            numpr.append(ni)
            insert_ordered(ppr, numpr, PPR_ORDER)
            self.add_inline(par, parse_inline(it["text"]))
            for c in it["children"]:
                if c["type"] == "quote":
                    self.add_quote(c["text"], 1.0 + 0.8 * level)
                else:
                    self.add_list(c, level + 1)
        # small gap after the top-level list
        if level == 0:
            d.paragraphs[-1].paragraph_format.space_after = Pt(8)

    def add_table(self, b):
        ncols = len(b["header"])
        weights = []
        for c in range(ncols):
            cells = [b["header"][c]] + [r[c] for r in b["rows"] if c < len(r)]
            lens = [len(plain(parse_inline(x))) for x in cells]
            weights.append(min(max(max(lens) * 0.7 + sum(lens) / len(lens) * 0.3, 6), 60))
        total = sum(weights)
        widths = [int(TEXT_WIDTH_DXA * w / total) for w in weights]
        widths[-1] += TEXT_WIDTH_DXA - sum(widths)

        table = self.doc.add_table(rows=1 + len(b["rows"]), cols=ncols)
        table.style = "Table Grid"
        table.autofit = False
        tblpr = table._tbl.tblPr
        for tag, attrs in (("w:tblW", {"w:w": str(TEXT_WIDTH_DXA), "w:type": "dxa"}),
                           ("w:tblLayout", {"w:type": "fixed"})):
            el = tblpr.find(qn(tag))
            if el is None:
                el = OxmlElement(tag)
                insert_ordered(tblpr, el, TBLPR_ORDER)
            for k, v in attrs.items():
                el.set(qn(k), v)
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            + "".join(
                f'<w:{s} w:val="single" w:sz="4" w:space="0" w:color="{GRAY_BORDER}"/>'
                for s in ("top", "left", "bottom", "right", "insideH", "insideV")
            )
            + "</w:tblBorders>"
        )
        old = tblpr.find(qn("w:tblBorders"))
        if old is not None:
            tblpr.remove(old)
        insert_ordered(tblpr, borders, TBLPR_ORDER)
        grid = table._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))

        all_rows = [b["header"]] + b["rows"]
        for ri, row in enumerate(all_rows):
            tr = table.rows[ri]
            trpr = tr._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if ri == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for ci in range(ncols):
                cell = tr.cells[ci]
                cell.width = Emu_from_dxa(widths[ci])
                tcpr = cell._tc.get_or_add_tcPr()
                w_el = tcpr.find(qn("w:tcW"))
                w_el.set(qn("w:w"), str(widths[ci]))
                w_el.set(qn("w:type"), "dxa")
                if ri == 0:
                    shade(tcpr, BRAND_TINT)
                text = row[ci] if ci < len(row) else ""
                par = cell.paragraphs[0]
                par.paragraph_format.space_after = Pt(2)
                par.paragraph_format.space_before = Pt(2)
                par.paragraph_format.line_spacing = 1.05
                self.add_inline(par, parse_inline(text), size=10, bold=(ri == 0))
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)

    def build(self, blocks, out: Path):
        # bookmarks for every heading first, so links can resolve forward
        for b in blocks:
            if b["type"] == "h":
                s = slug(plain(parse_inline(b["text"])))
                self.anchors[s] = f"sec_{len(self.anchors) + 1}"
        for b in blocks:
            self.add_block(b)
        self.doc.save(out)


def Emu_from_dxa(dxa):
    from docx.shared import Twips
    return Twips(dxa)


def main():
    for stem, title in GUIDES:
        md = (DOCS / f"{stem}.md").read_text(encoding="utf-8")
        blocks = parse_markdown(md)
        anchors = {slug(plain(parse_inline(b["text"]))) for b in blocks if b["type"] == "h"}
        header = f"Cổng Đặt Hàng – IMV · {title}"

        # warn about anchors that don't resolve
        for m in re.finditer(r"\]\(#([^)]+)\)", md):
            if m.group(1) not in anchors:
                print(f"  ! unresolved anchor in {stem}: #{m.group(1)}")

        DocxBuilder(title, header).build(blocks, DOCS / f"{stem}.docx")
        build_pdf(build_html(blocks, title, anchors), DOCS / f"{stem}.pdf", header)
        print(f"built {stem}.docx and {stem}.pdf")


if __name__ == "__main__":
    sys.exit(main())
